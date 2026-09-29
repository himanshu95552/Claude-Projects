import json, csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
S = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'private'); OUT = sys.argv[1]
d = json.load(open(S + '/merged.json'))
F = 'Arial'; HF = Font(name=F, bold=True, color='FFFFFF', size=10); BF = Font(name=F, size=10)
HFILL = PatternFill('solid', fgColor='1F3A5F'); LINK = Font(name=F, size=10, color='0563C1', underline='single')
wb = Workbook()

def sheet(title, cols, rows, widths=None, tname=None):
    ws = wb.create_sheet(title)
    ws.append(cols)
    for c in ws[1]: c.font = HF; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append([r.get(c, '') for c in cols])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BF; c.alignment = Alignment(vertical='top', wrap_text=False)
            v = c.value
            if isinstance(v, str) and v.startswith('http') and ' ' not in v:
                c.hyperlink = v; c.font = LINK
    for i, col in enumerate(cols, 1):
        w = (widths or {}).get(col) or min(max(len(col) + 2, max((len(str(r.get(col, ''))) for r in rows), default=10) + 2), 45)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'C2' if len(cols) > 3 else 'A2'
    ref = f'A1:{get_column_letter(len(cols))}{len(rows) + 1}'
    t = Table(displayName=tname or title.replace(' ', '').replace('&', '').replace('(', '').replace(')', '').replace('–', ''), ref=ref)
    t.tableStyleInfo = TableStyleInfo(name='TableStyleLight9', showRowStripes=True)
    ws.add_table(t)
    ws.row_dimensions[1].height = 30
    return ws

people_cols = d['people_cols']
wsP = sheet('People – Prospects', people_cols, d['people'], {'Priority': 9, 'Name': 28, 'Signal / evidence': 60, 'Title / role': 38, 'Source': 40}, 'People')
bcols = ['Company', 'Segment', 'Booth', 'Featured exhibitor', 'Website', 'LinkedIn', 'X / Twitter', 'Instagram', 'Facebook', 'YouTube',
         'Public contact email', 'Press / media email', 'HQ', 'What they sell', 'Key people (name - title - LinkedIn)', 'RSNA 2026 page / meeting form', 'MYS exhibitor profile', 'Fields to double-check']
sheet('Brands – Enriched', bcols, d['brands'], {'What they sell': 55, 'Key people (name - title - LinkedIn)': 70}, 'Brands')
excols = ['Exhibitor', 'Segment', 'Booth', 'Featured exhibitor', 'Enriched in Brands tab', 'MYS exhibitor profile']
exrows = [{'Exhibitor': r['exhibitor'], 'Segment': r['segment'], 'Booth': r['booth'], 'Featured exhibitor': r['featured'], 'Enriched in Brands tab': r.get('enriched', ''), 'MYS exhibitor profile': r['mys_profile_url']} for r in sorted(d['exhibitors'], key=lambda r: r['exhibitor'].lower())]
sheet('All Exhibitors (716)', excols, exrows, {'Exhibitor': 50, 'MYS exhibitor profile': 60}, 'Exhibitors')
sheet('Org & Press Contacts', ['Contact', 'Purpose', 'Email', 'Phone', 'Source'], d['orgc'], None, 'OrgContacts')
pb = list(csv.DictReader(open(S + '/raw/playbook.tsv'), delimiter='\t'))
sheet('Search Playbook', ['platform', 'goal', 'exact search string / filter', 'why it works'], pb, {'exact search string / filter': 80, 'why it works': 70}, 'Playbook')
tp = [{k: v.replace('\\n', '\n') for k, v in r.items()} for r in csv.DictReader(open(S + '/raw/templates.tsv'), delimiter='\t')]
ws = sheet('Outreach Templates', ['tag', 'channel', 'message template'], tp, {'message template': 110}, 'Templates')
for row in ws.iter_rows(min_row=2):
    row[2].alignment = Alignment(wrap_text=True, vertical='top')

# ---- Start Here ----
st = wb['Sheet']; st.title = 'Start Here'
st['A1'] = 'RSNA 2026 Prospect List – meetup outreach'; st['A1'].font = Font(name=F, bold=True, size=16, color='1F3A5F')
st['A2'] = 'RSNA 2026 · McCormick Place, Chicago · Nov 29 – Dec 3, 2026 (technical exhibits Nov 29 – Dec 2) · compiled 2026-09-29 from public sources'
st['A2'].font = Font(name=F, italic=True, size=10)
lines = [
 ('What is in this workbook', ''),
 ('People – Prospects', 'Named individuals, tagged by role (Authorized Official, Speaker, Influencer, Media / Press, Society Official, Brand Representative, Startup Founder, Health System Recruiter, Radiologist / Clinical Leader, Service Provider). Filter on "Primary tag" and "Priority".'),
 ('Brands – Enriched', 'Priority exhibitors with website, socials, published emails, HQ, what they sell, key decision-makers and RSNA 2026 meeting-request pages.'),
 ('All Exhibitors (716)', 'Every company on the official RSNA 2026 exhibitor list (Map Your Show, pulled 2026-09-29), auto-tagged by segment, booth where published.'),
 ('Org & Press Contacts', 'Official RSNA, society and trade-press inboxes/phones published on their own sites.'),
 ('Search Playbook', 'The exact keyword / hashtag / filter strings used, per platform, to keep growing the list every week until the show.'),
 ('Outreach Templates', 'Short first-touch messages per tag.'),
 ('', ''),
 ('How to read the columns', ''),
 ('Priority A', 'Confirmed for RSNA 2026 (official role, program, or own public post) AND a reachable handle or email is known. Contact first.'),
 ('Priority B', 'Either confirmed OR reachable, not both (e.g. exhibitor staff, likely-attending influencers).'),
 ('Priority C', 'Role-based prospect; confirm attendance with the first message.'),
 ('RSNA 2026 status', '"Confirmed" = public evidence it is 2026 (link in Evidence URL). "Exhibiting company" = company is on the official exhibitor list; the named person is a decision-maker there, not individually confirmed.'),
 ('Public email', 'Only emails published on an official/public page are included. Nothing was guessed or pattern-constructed. For everyone else, "Email status" tells you how to get it (LinkedIn URL -> Apollo / Hunter / Crustdata / RocketReach).'),
 ('Fields to double-check', 'Fields the research agent could not back with a primary source; confirm before sending.'),
 ('Compliance', 'Business outreach only. Follow CAN-SPAM (US) / CASL / GDPR (EU contacts): identify yourself, give an opt-out, do not add personal emails to bulk sequences, honour unsubscribes. Many RSNA attendees are EU/UK-based.'),
 ('', ''),
 ('Counts (live formulas)', ''),
]
r = 4
for a, b in lines:
    st.cell(r, 1, a).font = Font(name=F, bold=True, size=11 if not b else 10, color='1F3A5F' if not b else '000000')
    st.cell(r, 2, b).font = BF; st.cell(r, 2).alignment = Alignment(wrap_text=True, vertical='top')
    r += 1
n = len(d['people']) + 1
P = "'People – Prospects'"
tags = ['Authorized Official', 'Speaker', 'Influencer', 'Media / Press', 'Society Official', 'Brand Representative', 'Startup Founder', 'Radiologist / Clinical Leader', 'Health System Recruiter', 'Service Provider', 'Investor', 'Interested / Attending']
st.cell(r, 1, 'People (total)').font = BF; st.cell(r, 2, f"=COUNTA({P}!B2:B{n})").font = BF; r += 1
for pr in ['A', 'B', 'C']:
    st.cell(r, 1, f'  Priority {pr}').font = BF; st.cell(r, 2, f'=COUNTIF({P}!A2:A{n},"{pr}")').font = BF; r += 1
for t in tags:
    st.cell(r, 1, f'  Tag: {t}').font = BF; st.cell(r, 2, f'=COUNTIF({P}!C2:C{n},"{t}")').font = BF; r += 1
st.cell(r, 1, '  With a published email').font = BF; st.cell(r, 2, f'=COUNTIF({P}!P2:P{n},"?*")').font = BF; r += 1
st.cell(r, 1, '  With LinkedIn URL').font = BF; st.cell(r, 2, f'=COUNTIF({P}!L2:L{n},"?*")').font = BF; r += 1
ne = len(d['exhibitors']) + 1; E = "'All Exhibitors (716)'"
st.cell(r, 1, 'Exhibitors (total)').font = BF; st.cell(r, 2, f'=COUNTA({E}!A2:A{ne})').font = BF; r += 1
st.cell(r, 1, '  with booth number listed').font = BF; st.cell(r, 2, f'=COUNTIF({E}!C2:C{ne},"?*")').font = BF; r += 1
nb = len(d['brands']) + 1
st.cell(r, 1, 'Brands enriched').font = BF; st.cell(r, 2, f"=COUNTA('Brands – Enriched'!A2:A{nb})").font = BF; r += 1
st.column_dimensions['A'].width = 34; st.column_dimensions['B'].width = 120
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)  # LibreOffice recalc unavailable in the build sandbox; Excel/Sheets compute on open
wb.save(OUT)
print('saved', OUT)
