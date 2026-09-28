"""Build the Alpha Nodus social media performance & demo-attribution tracker.

Run:  python build_tracker.py  ->  AlphaNodus_Social_Media_Tracker.xlsx
Then recalculate with LibreOffice (xlsx skill scripts/recalc.py) so cached
values exist for previewers.
"""
import datetime as dt
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("AlphaNodus_Social_Media_Tracker.xlsx")

FONT = "Arial"
NAVY, TEAL, GREEN, ORANGE, PURPLE, GREY = "1F3864", "1F6F78", "548235", "C55A11", "7030A0", "595959"
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")   # light yellow = you type here
CALC_FILL = PatternFill("solid", fgColor="F2F2F2")    # light grey = formula, don't touch
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PT_FIRST, PT_LAST = 5, 1004      # Post Tracker data rows (1,000 posts)
DP_FIRST, DP_LAST = 5, 504       # Demo Pipeline rows (500 demos)
CS_FIRST, CS_LAST = 5, 54        # Case Studies rows (50)

FMT_DATE = "mm/dd/yyyy"
FMT_TIME = "h:mm AM/PM"
FMT_INT = "#,##0"
FMT_PCT = "0.0%"
FMT_USD = "$#,##0"
FMT_X = '0.00"x"'


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = f(True, NAVY, 16)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = f(False, GREY, 10, True)


def add_list_validation(ws, rng, source, prompt=None):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    dv.error = "Pick a value from the dropdown (lists live on the Settings sheet)."
    dv.errorTitle = "Invalid entry"
    dv.showErrorMessage = True
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


# ---------------------------------------------------------------- Settings
LISTS = {
    "Platform": ["LinkedIn", "Instagram", "X", "Facebook", "YouTube"],
    "Platform Code": ["LI", "IG", "X", "FB", "YT"],
    "utm_source": ["linkedin", "instagram", "twitter", "facebook", "youtube"],
    "Account": ["Company Page", "Founder / Exec Profile", "Employee Advocate"],
    "Account Code": ["C", "F", "E"],
    "Content Pillar": [
        "Case Study / Customer Result", "Product Demo / Feature", "Industry Insight / Data",
        "Thought Leadership / POV", "Customer Testimonial", "Reimbursement & Regulatory",
        "Operations Tips (Scheduling, No-shows, Throughput)", "Team / Culture / Behind the Scenes",
        "Event / Webinar / Conference", "Offer / Direct Demo CTA",
    ],
    "Pillar Slug": [
        "case-study", "product", "insight", "thought-leadership", "testimonial", "regulatory",
        "ops-tips", "culture", "event", "demo-offer",
    ],
    "Post Format": [
        "Text Only", "Single Image", "Carousel / Document (PDF)", "Short Video / Reel", "Long Video",
        "Story", "Poll", "Thread", "Article / Newsletter", "Link Post", "Live / Event",
    ],
    "Hook Type": [
        "Stat / Number", "Question", "Story / Anecdote", "Contrarian Take", "How-To / List",
        "Customer Quote", "News Hook", "Before / After",
    ],
    "Target Persona": [
        "Owner / CEO", "Medical Director / Radiologist", "Practice / Operations Manager",
        "Revenue Cycle / Billing Lead", "IT / PACS Administrator", "Front Desk / Scheduling Lead",
        "General / Mixed",
    ],
    "CTA": ["Book a Demo", "Visit Website", "Read Case Study", "Download Resource", "Comment / Engage",
            "Follow", "No CTA"],
    "Paid Status": ["Organic", "Boosted", "Paid Ad"],
    "Link Placement": ["In Post Body", "First Comment", "Link in Bio", "Story Link Sticker",
                       "Link Preview Card", "No Link"],
    "Next Action": ["Scale It (repeat + boost)", "Iterate (new hook/format)", "Retire", "Too Early to Tell"],
    "Demo Status": ["Scheduled", "Attended", "No-Show", "Rescheduled", "Cancelled"],
    "Qualified?": ["Yes", "No", "TBD"],
    "Deal Stage": ["Demo Booked", "Demo Completed", "Proposal Sent", "Negotiation", "Closed Won",
                   "Closed Lost", "Nurture"],
    "Lead Source": ["LinkedIn", "Instagram", "X", "Facebook", "YouTube", "Google Search", "Direct / Typed URL",
                    "Email", "Referral", "Event / Conference", "Social DM", "Other"],
    "How Heard": ["LinkedIn - company post", "LinkedIn - founder/team post", "Instagram", "X", "Facebook",
                  "YouTube", "Google search", "Colleague / referral", "Conference / event", "Email",
                  "Other"],
    "Client Type": ["Independent Imaging Center", "Multi-site Imaging Network", "Radiology Group / Practice",
                    "Outpatient / ASC Imaging", "Teleradiology", "Other"],
    "Proof Type": ["Hard Metric", "Customer Quote", "Video Testimonial", "Before / After", "Screenshot / Demo"],
    "Day": ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"],
    "Checklist Status": ["Not Started", "In Progress", "Done", "N/A"],
    "US State": ["AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "DC", "FL", "GA", "HI", "ID", "IL", "IN",
                 "IA", "KS", "KY", "LA", "ME", "MD", "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH",
                 "NJ", "NM", "NY", "NC", "ND", "OH", "OK", "OR", "PA", "RI", "SC", "SD", "TN", "TX", "UT",
                 "VT", "VA", "WA", "WV", "WI", "WY"],
}

wb = Workbook()
ws_start = wb.active
ws_start.title = "Start Here"
ws_pt = wb.create_sheet("Post Tracker")
ws_dp = wb.create_sheet("Demo Pipeline")
ws_db = wb.create_sheet("Dashboard")
ws_cs = wb.create_sheet("Case Studies")
ws_fg = wb.create_sheet("Follower Growth")
ws_ut = wb.create_sheet("Tracking Setup")
ws_md = wb.create_sheet("Metric Definitions")
ws_set = wb.create_sheet("Settings")

# Settings: key settings block (A:B), then each list in its own column from D.
title(ws_set, "Settings & Dropdown Lists",
      "Yellow cells are editable. Add items to the bottom of a list and every dropdown picks them up.")
ws_set["A4"], ws_set["B4"] = "Setting", "Value"
settings = [
    ("Default landing page (used when a post's Landing Page is blank)", "https://www.alphanodus.com/"),
    ("Engagement weight: Like / Reaction", 1),
    ("Engagement weight: Comment", 3),
    ("Engagement weight: Share / Repost", 4),
    ("Engagement weight: Save / Bookmark", 3),
    ("Engagement weight: Link Click", 2),
]
SET = {}
for i, (k, v) in enumerate(settings):
    r = 5 + i
    ws_set.cell(r, 1, k).font = f()
    c = ws_set.cell(r, 2, v)
    c.font, c.fill, c.border = f(color="0000FF"), INPUT_FILL, BORDER
    SET[k] = f"Settings!$B${r}"
DEFAULT_LP = SET["Default landing page (used when a post's Landing Page is blank)"]
ws_set["A12"] = ("Weights are a team judgement call, not an industry standard: comments and shares signal more "
                 "intent than a like, so they count more. Change them if your team disagrees.")
ws_set["A12"].font = f(italic=True, color=GREY)
ws_set["A12"].alignment = Alignment(wrap_text=True, vertical="top")
ws_set.row_dimensions[12].height = 60
for cell in ("A4", "B4"):
    ws_set[cell].font, ws_set[cell].fill = f(True, "FFFFFF"), fill(NAVY)
ws_set.column_dimensions["A"].width = 44
ws_set.column_dimensions["B"].width = 30

LIST_REF = {}      # name -> absolute range (with 40 spare slots per list)
LIST_COL = {}      # name -> column letter on Settings
col = 4
for name, items in LISTS.items():
    L = get_column_letter(col)
    h = ws_set.cell(4, col, name)
    h.font, h.fill, h.alignment = f(True, "FFFFFF"), fill(TEAL), Alignment(wrap_text=True)
    for i, it in enumerate(items):
        c = ws_set.cell(5 + i, col, it)
        c.font, c.fill, c.border = f(), INPUT_FILL, BORDER
    LIST_COL[name] = L
    LIST_REF[name] = f"Settings!${L}$5:${L}${5 + max(len(items), 1) + 39}"
    ws_set.column_dimensions[L].width = max(14, min(34, max(len(x) for x in items) + 2))
    col += 1
ws_set.freeze_panes = "A5"


# ------------------------------------------------------------ Post Tracker
# (key, header, kind, width, number_format, comment)  kind: in | calc
GROUPS = [
    ("POST DETAILS", NAVY, [
        ("post_id", "Post ID", "calc", 17, None, "Auto: platform + account code, date (yymmdd), time (hhmm). Used as utm_content — this is how demos are matched back to the post."),
        ("date", "Publish Date", "in", 12, FMT_DATE, None),
        ("time", "Publish Time", "in", 11, FMT_TIME, "Local time. Makes the Post ID unique when you post twice a day."),
        ("day", "Day", "calc", 6, None, None),
        ("week", "Week Of (Mon)", "calc", 12, FMT_DATE, None),
        ("month", "Month", "calc", 9, None, None),
        ("platform", "Platform", "in", 11, None, None),
        ("account", "Account", "in", 20, None, "Company Page vs Founder / Exec profile vs Employee reshare — founder posts usually behave very differently on LinkedIn, so track them separately."),
        ("format", "Post Format", "in", 22, None, None),
        ("pillar", "Content Pillar", "in", 28, None, None),
        ("case_study", "Case Study Featured", "in", 28, None, "Dropdown comes from the Case Studies sheet. Leave blank if the post doesn't feature one."),
        ("topic", "Topic / Headline / First Line", "in", 40, None, None),
        ("hook", "Hook Type", "in", 16, None, None),
        ("persona", "Target Persona", "in", 24, None, None),
        ("cta", "CTA", "in", 16, None, None),
        ("paid", "Organic / Paid", "in", 11, None, None),
        ("spend", "Ad Spend ($)", "in", 10, FMT_USD, None),
        ("post_url", "Live Post URL", "in", 34, None, None),
        ("followers", "Followers at Posting", "in", 11, FMT_INT, "Account follower count on the day of posting — needed for Reach % of Followers."),
    ]),
    ("TRACKING LINK (UTM)", TEAL, [
        ("landing", "Landing Page URL", "in", 30, None, "Page the post links to. Leave blank to use the default on the Settings sheet."),
        ("placement", "Link Placement", "in", 16, None, "Instagram captions aren't clickable, so note where the link lived (bio, story sticker…)."),
        ("utm_source", "utm_source", "calc", 11, None, None),
        ("utm_medium", "utm_medium", "calc", 11, None, None),
        ("utm_campaign", "utm_campaign", "calc", 20, None, "Case study slug if one is featured, otherwise the content-pillar slug."),
        ("utm_content", "utm_content", "calc", 17, None, None),
        ("tracking_url", "Tracking URL (copy into post)", "calc", 60, None, "Unique link for this post. Copy-paste it (or its short link) into the post."),
        ("short_link", "Short Link (optional)", "in", 22, None, "bit.ly / Dub / branded short link pointing at the Tracking URL."),
    ]),
    ("PLATFORM METRICS  (enter ~7 days after posting)", PURPLE, [
        ("pulled", "Data Pulled On", "in", 12, FMT_DATE, "Pull every post at the same age (e.g. day 7) so posts are comparable."),
        ("impr", "Impressions / Views", "in", 12, FMT_INT, "LinkedIn: Impressions. Instagram: Views. X: Impressions."),
        ("reach", "Reach (Reported)", "in", 11, FMT_INT, "Unique accounts/members reached, if the platform shows it (LinkedIn page posts: Members reached; Instagram: Accounts reached). X doesn't report it — leave blank."),
        ("video_views", "Video Views", "in", 10, FMT_INT, None),
        ("likes", "Likes / Reactions", "in", 10, FMT_INT, None),
        ("comments", "Comments", "in", 10, FMT_INT, None),
        ("shares", "Shares / Reposts", "in", 10, FMT_INT, None),
        ("saves", "Saves / Bookmarks", "in", 10, FMT_INT, None),
        ("clicks", "Link Clicks", "in", 10, FMT_INT, "Platform-reported clicks (or short-link clicks if the platform doesn't show them)."),
        ("profile_visits", "Profile Visits", "in", 10, FMT_INT, None),
        ("new_followers", "New Followers", "in", 10, FMT_INT, None),
        ("icp", "ICP Engagers", "in", 10, FMT_INT, "How many of the people who liked/commented/reposted are imaging-center or radiology decision-makers. Skim the engager list — this is the best early signal you're reaching buyers, not just peers."),
        ("dms", "Inbound DMs / Conversations", "in", 12, FMT_INT, None),
    ]),
    ("REACH & ENGAGEMENT (calculated)", GREEN, [
        ("reach_calc", "Reach (Calc.)", "calc", 11, FMT_INT, "Reported reach; if the platform doesn't report it, falls back to impressions."),
        ("reach_basis", "Reach Basis", "calc", 18, None, None),
        ("reach_rate", "Reach % of Followers", "calc", 11, FMT_PCT, ">100% means the post travelled beyond your followers."),
        ("eng", "Total Engagements", "calc", 11, FMT_INT, "Likes + Comments + Shares + Saves + Link Clicks"),
        ("er", "Engagement Rate (by Reach)", "calc", 12, FMT_PCT, "Primary metric: Total Engagements ÷ Reach (Calc.)"),
        ("er_impr", "Engagement Rate (by Impr.)", "calc", 12, FMT_PCT, None),
        ("ctr", "Click-Through Rate", "calc", 10, FMT_PCT, "Link Clicks ÷ Impressions"),
        ("amp", "Share Rate", "calc", 10, FMT_PCT, "Shares ÷ Reach — how often people pass it on."),
        ("wscore", "Weighted Eng. per 1K Reach", "calc", 12, "0.0", "Weighted engagements (weights on Settings) per 1,000 people reached."),
        ("vs_avg", "ER vs Platform Avg", "calc", 11, FMT_X, "This post's Engagement Rate ÷ the average for its platform. 1.2x+ = green, 0.8x or less = red."),
    ]),
    ("WEBSITE & DEMO FUNNEL", ORANGE, [
        ("sessions", "Website Sessions (GA4)", "in", 11, FMT_INT, "GA4 sessions where Session manual ad content = this Post ID (lower-case). See Tracking Setup."),
        ("engaged", "Engaged Sessions (GA4)", "in", 11, FMT_INT, None),
        ("demo_page", "Demo Page Visits", "in", 11, FMT_INT, "Sessions from this post that reached the demo / book-a-demo page."),
        ("form_starts", "Demo Form Starts (optional)", "in", 11, FMT_INT, None),
        ("booked", "Demos Booked", "calc", 10, FMT_INT, "Auto-counted from the Demo Pipeline sheet (rows whose Attributed Post ID = this Post ID)."),
        ("attended", "Demos Attended", "calc", 10, FMT_INT, None),
        ("qualified", "Qualified Demos", "calc", 10, FMT_INT, None),
        ("won", "Deals Won", "calc", 9, FMT_INT, None),
        ("pipeline", "Pipeline Value ($)", "calc", 12, FMT_USD, None),
        ("click_sess", "Click → Session %", "calc", 10, FMT_PCT, "Low values = bot clicks, slow page, or link-preview clicks that never loaded the site."),
        ("sess_demo", "Session → Demo Page %", "calc", 10, FMT_PCT, None),
        ("demo_book", "Demo Page → Booked %", "calc", 10, FMT_PCT, None),
        ("sess_book", "Session → Booked %", "calc", 10, FMT_PCT, None),
        ("cpd", "Cost per Demo ($)", "calc", 10, FMT_USD, None),
    ]),
    ("LEARNINGS", GREY, [
        ("learning", "What We Learned", "in", 40, None, None),
        ("next", "Next Action", "in", 22, None, None),
    ]),
    ("HELPERS (Dashboard)", "A6A6A6", [
        ("rk_er", "Rank Key: ER", "calc", 10, "0.000", None),
        ("rk_demo", "Rank Key: Demos", "calc", 10, "0.000", None),
    ]),
]

C = {}          # key -> column letter on Post Tracker
col = 1
for _, _, cols in GROUPS:
    for spec in cols:
        C[spec[0]] = get_column_letter(col)
        col += 1
PT_LASTCOL = col - 1


def rng(key, sheet="'Post Tracker'!"):
    return f"{sheet}${C[key]}${PT_FIRST}:${C[key]}${PT_LAST}"


def pt_formula(key, r):
    x = {k: f"{v}{r}" for k, v in C.items()}
    plat_name, plat_code, utm_src = LIST_REF["Platform"], LIST_REF["Platform Code"], LIST_REF["utm_source"]
    acct_name, acct_code = LIST_REF["Account"], LIST_REF["Account Code"]
    pil, pil_slug = LIST_REF["Content Pillar"], LIST_REF["Pillar Slug"]
    cs_name = f"'Case Studies'!$B${CS_FIRST}:$B${CS_LAST}"
    cs_slug = f"'Case Studies'!$C${CS_FIRST}:$C${CS_LAST}"
    dp = "'Demo Pipeline'!"
    dp_id = f"{dp}$Q${DP_FIRST}:$Q${DP_LAST}"
    w = [SET[f"Engagement weight: {n}"] for n in ("Like / Reaction", "Comment", "Share / Repost",
                                                    "Save / Bookmark", "Link Click")]
    pid = x["post_id"]
    F = {
        "post_id": (f'=IF(OR({x["date"]}="",{x["platform"]}="",{x["account"]}=""),"",'
                    f'IFERROR(INDEX({plat_code},MATCH({x["platform"]},{plat_name},0))'
                    f'&INDEX({acct_code},MATCH({x["account"]},{acct_name},0))'
                    f'&"-"&TEXT({x["date"]},"yymmdd")&"-"&IF({x["time"]}="","0000",TEXT({x["time"]},"hhmm")),"CHECK LISTS"))'),
        "day": f'=IF({x["date"]}="","",TEXT({x["date"]},"ddd"))',
        "week": f'=IF({x["date"]}="","",{x["date"]}-WEEKDAY({x["date"]},3))',
        "month": f'=IF({x["date"]}="","",TEXT({x["date"]},"yyyy-mm"))',
        "utm_source": f'=IF({pid}="","",IFERROR(INDEX({utm_src},MATCH({x["platform"]},{plat_name},0)),""))',
        "utm_medium": f'=IF({pid}="","",IF(OR({x["paid"]}="",{x["paid"]}="Organic"),"social","paid_social"))',
        "utm_campaign": (f'=IF({pid}="","",IF({x["case_study"]}<>"",IFERROR(INDEX({cs_slug},MATCH({x["case_study"]},{cs_name},0))&"","case-study"),'
                         f'IF({x["pillar"]}<>"",IFERROR(INDEX({pil_slug},MATCH({x["pillar"]},{pil},0)),"general"),"general")))'),
        "utm_content": f'=IF({pid}="","",LOWER({pid}))',
        "tracking_url": (f'=IF({pid}="","",IF({x["landing"]}="",{DEFAULT_LP},{x["landing"]})'
                         f'&IF(ISNUMBER(FIND("?",IF({x["landing"]}="",{DEFAULT_LP},{x["landing"]}))),"&","?")'
                         f'&"utm_source="&{x["utm_source"]}&"&utm_medium="&{x["utm_medium"]}'
                         f'&"&utm_campaign="&{x["utm_campaign"]}&"&utm_content="&{x["utm_content"]})'),
        "reach_calc": f'=IF({pid}="","",IF({x["reach"]}<>"",{x["reach"]},IF({x["impr"]}<>"",{x["impr"]},"")))',
        "reach_basis": f'=IF({x["reach_calc"]}="","",IF({x["reach"]}<>"","Reported reach","Impressions (proxy)"))',
        "reach_rate": f'=IF(OR({x["reach_calc"]}="",N({x["followers"]})=0),"",{x["reach_calc"]}/{x["followers"]})',
        "eng": (f'=IF({pid}="","",IF(COUNT({x["likes"]}:{x["clicks"]})=0,"",'
                f'SUM({x["likes"]}:{x["clicks"]})))'),
        "er": f'=IF(OR({x["eng"]}="",N({x["reach_calc"]})=0),"",{x["eng"]}/{x["reach_calc"]})',
        "er_impr": f'=IF(OR({x["eng"]}="",N({x["impr"]})=0),"",{x["eng"]}/{x["impr"]})',
        "ctr": f'=IF(OR({x["clicks"]}="",N({x["impr"]})=0),"",{x["clicks"]}/{x["impr"]})',
        "amp": f'=IF(OR({x["shares"]}="",N({x["reach_calc"]})=0),"",{x["shares"]}/{x["reach_calc"]})',
        "wscore": (f'=IF(OR({x["eng"]}="",N({x["reach_calc"]})=0),"",'
                   f'(N({x["likes"]})*{w[0]}+N({x["comments"]})*{w[1]}+N({x["shares"]})*{w[2]}'
                   f'+N({x["saves"]})*{w[3]}+N({x["clicks"]})*{w[4]})/{x["reach_calc"]}*1000)'),
        "vs_avg": f'=IF({x["er"]}="","",IFERROR({x["er"]}/AVERAGEIFS({rng("er")},{rng("platform")},{x["platform"]}),""))',
        "booked": f'=IF({pid}="","",COUNTIFS({dp_id},{pid}))',
        "attended": f'=IF({pid}="","",COUNTIFS({dp_id},{pid},{dp}$W${DP_FIRST}:$W${DP_LAST},"Attended"))',
        "qualified": f'=IF({pid}="","",COUNTIFS({dp_id},{pid},{dp}$X${DP_FIRST}:$X${DP_LAST},"Yes"))',
        "won": f'=IF({pid}="","",COUNTIFS({dp_id},{pid},{dp}$Y${DP_FIRST}:$Y${DP_LAST},"Closed Won"))',
        "pipeline": f'=IF({pid}="","",SUMIFS({dp}$Z${DP_FIRST}:$Z${DP_LAST},{dp_id},{pid}))',
        "click_sess": f'=IF(OR({x["sessions"]}="",N({x["clicks"]})=0),"",{x["sessions"]}/{x["clicks"]})',
        "sess_demo": f'=IF(OR({x["demo_page"]}="",N({x["sessions"]})=0),"",{x["demo_page"]}/{x["sessions"]})',
        "demo_book": f'=IF(OR({pid}="",N({x["demo_page"]})=0),"",{x["booked"]}/{x["demo_page"]})',
        "sess_book": f'=IF(OR({pid}="",N({x["sessions"]})=0),"",{x["booked"]}/{x["sessions"]})',
        "cpd": f'=IF(OR(N({x["spend"]})=0,N({x["booked"]})=0),"",{x["spend"]}/{x["booked"]})',
        "rk_er": (f'=IF(AND(ISNUMBER({x["date"]}),{x["date"]}>=Dashboard!$C$4,{x["date"]}<=Dashboard!$F$4,'
                  f'ISNUMBER({x["er"]})),{x["er"]}+ROW()/10^9,"")'),
        "rk_demo": (f'=IF(AND(ISNUMBER({x["date"]}),{x["date"]}>=Dashboard!$C$4,{x["date"]}<=Dashboard!$F$4,{pid}<>""),'
                    f'N({x["booked"]})+N({x["er"]})/100+ROW()/10^9,"")'),
    }
    return F[key]


title(ws_pt, "Post Tracker — one row per post, per platform",
      "Yellow = you fill in · Grey = calculated, don't type over · Row 5 is an EXAMPLE — overwrite or delete it (and its Demo Pipeline row).")
col = 1
for gname, gcolor, cols in GROUPS:
    start = col
    for key, header, kind, width, nf, comment in cols:
        L = get_column_letter(col)
        h = ws_pt.cell(4, col, header)
        h.font = f(True, "FFFFFF")
        h.fill = fill(gcolor if kind == "calc" else NAVY) if gname != "POST DETAILS" else fill(NAVY if kind == "in" else "44546A")
        if kind == "in" and gname != "POST DETAILS":
            h.fill = fill("2F5597")
        if kind == "calc":
            h.fill = fill(gcolor if gname != "POST DETAILS" else "44546A")
        h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        h.border = BORDER
        if comment:
            h.comment = Comment(comment, "Tracker")
        ws_pt.column_dimensions[L].width = width
        for r in range(PT_FIRST, PT_LAST + 1):
            c = ws_pt.cell(r, col)
            c.font = f()
            c.border = BORDER
            c.fill = CALC_FILL if kind == "calc" else INPUT_FILL
            if nf:
                c.number_format = nf
            if kind == "calc":
                c.value = pt_formula(key, r)
        col += 1
    ws_pt.merge_cells(start_row=3, start_column=start, end_row=3, end_column=col - 1)
    g = ws_pt.cell(3, start, gname)
    g.font, g.fill = f(True, "FFFFFF"), fill(gcolor)
    g.alignment = Alignment(horizontal="center")
ws_pt.row_dimensions[4].height = 45
ws_pt.freeze_panes = "C5"
ws_pt.auto_filter.ref = f"A4:{get_column_letter(PT_LASTCOL)}{PT_LAST}"
for key in ("rk_er", "rk_demo"):
    ws_pt.column_dimensions[C[key]].hidden = True

pt_dropdowns = {
    "platform": "Platform", "account": "Account", "format": "Post Format", "pillar": "Content Pillar",
    "hook": "Hook Type", "persona": "Target Persona", "cta": "CTA", "paid": "Paid Status",
    "placement": "Link Placement", "next": "Next Action",
}
for key, lst in pt_dropdowns.items():
    add_list_validation(ws_pt, f"{C[key]}{PT_FIRST}:{C[key]}{PT_LAST}", f"={LIST_REF[lst]}")
add_list_validation(ws_pt, f"{C['case_study']}{PT_FIRST}:{C['case_study']}{PT_LAST}",
                    f"='Case Studies'!$B${CS_FIRST}:$B${CS_LAST}")

# Conditional formatting
full = lambda k: f"{C[k]}{PT_FIRST}:{C[k]}{PT_LAST}"
ws_pt.conditional_formatting.add(
    full("post_id"),
    FormulaRule(formula=[f'AND(${C["post_id"]}{PT_FIRST}<>"",COUNTIF(${C["post_id"]}${PT_FIRST}:${C["post_id"]}${PT_LAST},${C["post_id"]}{PT_FIRST})>1)'],
                fill=fill("F8CBAD"), font=Font(name=FONT, color="9C0006", bold=True)))
ws_pt.conditional_formatting.add(full("vs_avg"), CellIsRule(operator="greaterThanOrEqual", formula=["1.2"], fill=fill("C6EFCE"), font=Font(name=FONT, color="006100", bold=True)))
ws_pt.conditional_formatting.add(full("vs_avg"), CellIsRule(operator="between", formula=["0.0001", "0.8"], fill=fill("FFC7CE"), font=Font(name=FONT, color="9C0006")))
ws_pt.conditional_formatting.add(full("er"), ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"))
ws_pt.conditional_formatting.add(full("booked"), CellIsRule(operator="greaterThan", formula=["0"], fill=fill("C6EFCE"), font=Font(name=FONT, color="006100", bold=True)))

# Example row (row 5)
example = {
    "date": dt.date(2026, 9, 22), "time": dt.time(9, 0), "platform": "LinkedIn", "account": "Company Page",
    "format": "Carousel / Document (PDF)", "pillar": "Case Study / Customer Result",
    "case_study": "EXAMPLE - MRI no-show reduction",
    "topic": "EXAMPLE - How a 3-site imaging center in Texas cut MRI no-shows in 90 days",
    "hook": "Stat / Number", "persona": "Owner / CEO", "cta": "Book a Demo", "paid": "Organic",
    "post_url": "https://www.linkedin.com/feed/update/urn:li:activity:EXAMPLE", "followers": 1250,
    "placement": "First Comment", "pulled": dt.date(2026, 9, 29), "impr": 2400, "reach": 1650,
    "likes": 48, "comments": 9, "shares": 5, "clicks": 37, "profile_visits": 22, "new_followers": 6,
    "icp": 7, "dms": 1, "sessions": 29, "engaged": 21, "demo_page": 8, "form_starts": 3,
    "learning": "EXAMPLE ROW - overwrite. Result stat on slide 1 + 'book a demo' in first comment drove clicks.",
    "next": "Scale It (repeat + boost)",
}
for k, v in example.items():
    c = ws_pt[f"{C[k]}{PT_FIRST}"]
    c.value = v
    c.font = f(italic=True, color="0000FF")


# ------------------------------------------------------------ Demo Pipeline
title(ws_dp, "Demo Pipeline — one row per demo booked (all sources, not only social)",
      "Copy the utm_* values from the demo form's hidden fields / CRM. No UTM? Put the Post ID in 'Manual Post ID' if the lead told you which post. Row 5 is an EXAMPLE.")
DP_COLS = [
    # (header, kind, width, fmt, dropdown/list, comment)
    ("Booking #", "calc", 9, '"D-"000', None, None),
    ("Date Booked", "in", 12, FMT_DATE, None, None),
    ("Contact Name", "in", 20, None, None, None),
    ("Role / Persona", "in", 22, None, "Target Persona", None),
    ("Imaging Center / Organization", "in", 28, None, None, None),
    ("City", "in", 14, None, None, None),
    ("State", "in", 7, None, "US State", None),
    ("# of Locations", "in", 9, FMT_INT, None, None),
    ("Modalities (MRI, CT, X-ray…)", "in", 18, None, None, None),
    ("Lead Source (Channel)", "in", 16, None, "Lead Source", None),
    ("utm_source", "in", 11, None, None, "From hidden form field / CRM."),
    ("utm_medium", "in", 11, None, None, None),
    ("utm_campaign", "in", 18, None, None, None),
    ("utm_content", "in", 17, None, None, "This is the Post ID (lower-case). It links the demo back to the post automatically."),
    ("How Did You Hear About Us? (self-reported)", "in", 24, None, "How Heard", "Ask this on the demo form. It catches 'dark social' — people who saw a post but typed the URL instead of clicking."),
    ("Manual Post ID (if no UTM)", "in", 17, None, None, "If the lead says 'I saw your post about X', find that Post ID on the Post Tracker and paste it here."),
    ("Attributed Post ID", "calc", 17, None, None, "utm_content if present, otherwise Manual Post ID."),
    ("Attribution Type", "calc", 18, None, None, None),
    ("Post Platform", "calc", 11, None, None, None),
    ("Post Case Study", "calc", 24, None, None, None),
    ("Days: Post → Booking", "calc", 10, FMT_INT, None, "How long the post took to turn into a demo — tells you how long to wait before judging a post."),
    ("Demo Date", "in", 12, FMT_DATE, None, None),
    ("Demo Status", "in", 12, None, "Demo Status", None),
    ("Qualified?", "in", 10, None, "Qualified?", None),
    ("Deal Stage", "in", 15, None, "Deal Stage", None),
    ("Est. Deal Value ($)", "in", 12, FMT_USD, None, None),
    ("CRM Record Link", "in", 24, None, None, None),
    ("Notes", "in", 30, None, None, None),
]
pt_id, pt_plat, pt_cs, pt_date = rng("post_id"), rng("platform"), rng("case_study"), rng("date")
for i, (hdr, kind, width, nf, lst, comment) in enumerate(DP_COLS, start=1):
    L = get_column_letter(i)
    h = ws_dp.cell(4, i, hdr)
    h.font = f(True, "FFFFFF")
    h.fill = fill("2F5597" if kind == "in" else ORANGE)
    h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    h.border = BORDER
    if comment:
        h.comment = Comment(comment, "Tracker")
    ws_dp.column_dimensions[L].width = width
    for r in range(DP_FIRST, DP_LAST + 1):
        c = ws_dp.cell(r, i)
        c.font, c.border = f(), BORDER
        c.fill = CALC_FILL if kind == "calc" else INPUT_FILL
        if nf:
            c.number_format = nf
        if kind == "calc":
            c.value = {
                "A": f'=IF(B{r}="","",ROW()-{DP_FIRST - 1})',
                "Q": f'=IF(N{r}<>"",UPPER(TRIM(N{r})),IF(P{r}<>"",UPPER(TRIM(P{r})),""))',
                "R": f'=IF(B{r}="","",IF(N{r}<>"","UTM tracked",IF(P{r}<>"","Manual / self-reported",IF(O{r}<>"","Self-reported channel only","Unattributed"))))',
                "S": f'=IF(Q{r}="","",IFERROR(INDEX({pt_plat},MATCH(Q{r},{pt_id},0)),"ID NOT FOUND"))',
                "T": f'=IF(Q{r}="","",IFERROR(INDEX({pt_cs},MATCH(Q{r},{pt_id},0))&"",""))',
                "U": f'=IF(OR(Q{r}="",B{r}=""),"",IFERROR(B{r}-INDEX({pt_date},MATCH(Q{r},{pt_id},0)),""))',
            }[L]
    if lst:
        add_list_validation(ws_dp, f"{L}{DP_FIRST}:{L}{DP_LAST}", f"={LIST_REF[lst]}")
ws_dp.row_dimensions[4].height = 45
ws_dp.freeze_panes = "C5"
ws_dp.auto_filter.ref = f"A4:{get_column_letter(len(DP_COLS))}{DP_LAST}"
ws_dp.conditional_formatting.add(f"S{DP_FIRST}:S{DP_LAST}", CellIsRule(operator="equal", formula=['"ID NOT FOUND"'], fill=fill("FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
dp_example = {
    "B": dt.date(2026, 9, 24), "C": "EXAMPLE - Jane Doe", "D": "Owner / CEO", "E": "EXAMPLE Imaging Partners",
    "F": "Dallas", "G": "TX", "H": 3, "I": "MRI, CT, X-ray", "J": "LinkedIn", "K": "linkedin", "L": "social",
    "M": "cs01-mri-noshows", "N": "lic-260922-0900", "O": "LinkedIn - company post", "V": dt.date(2026, 9, 30),
    "W": "Scheduled", "X": "TBD", "Y": "Demo Booked", "Z": 18000,
    "AB": "EXAMPLE ROW - delete. Deal value is illustrative only.",
}
for k, v in dp_example.items():
    c = ws_dp[f"{k}{DP_FIRST}"]
    c.value = v
    c.font = f(italic=True, color="0000FF")


# ------------------------------------------------------------ Case Studies
title(ws_cs, "Case Study Library — what each story is, and how it performs",
      "Add each case study once. Its name feeds the Post Tracker dropdown; performance columns fill automatically.")
CS_COLS = [
    ("CS ID", "in", 7, None, None),
    ("Case Study Name (short — shows in dropdown)", "in", 30, None, None),
    ("UTM Campaign Slug", "in", 20, None, "lower-case, hyphens, no spaces — e.g. cs02-scheduling-throughput"),
    ("Client Type", "in", 22, None, "Client Type"),
    ("Client State", "in", 8, None, "US State"),
    ("Modalities", "in", 14, None, None),
    ("Core Problem / Pain Point", "in", 30, None, None),
    ("Solution / Feature Showcased", "in", 28, None, None),
    ("Headline Result (the number)", "in", 26, None, None),
    ("Proof Type", "in", 16, None, "Proof Type"),
    ("Asset / Landing Page Link", "in", 26, None, None),
    ("# Posts", "calc", 8, FMT_INT, None),
    ("Total Reach", "calc", 10, FMT_INT, None),
    ("Engagements", "calc", 11, FMT_INT, None),
    ("Eng. Rate", "calc", 9, FMT_PCT, None),
    ("Link Clicks", "calc", 9, FMT_INT, None),
    ("Sessions", "calc", 9, FMT_INT, None),
    ("Demos Booked", "calc", 9, FMT_INT, None),
    ("Qualified Demos", "calc", 9, FMT_INT, None),
    ("Pipeline ($)", "calc", 11, FMT_USD, None),
]
for i, (hdr, kind, width, nf, lst) in enumerate(CS_COLS, start=1):
    L = get_column_letter(i)
    h = ws_cs.cell(4, i, hdr)
    h.font = f(True, "FFFFFF")
    h.fill = fill("2F5597" if kind == "in" else GREEN)
    h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    h.border = BORDER
    if lst == "lower-case, hyphens, no spaces — e.g. cs02-scheduling-throughput":
        h.comment = Comment(lst, "Tracker")
        lst = None
    ws_cs.column_dimensions[L].width = width
    for r in range(CS_FIRST, CS_LAST + 1):
        c = ws_cs.cell(r, i)
        c.font, c.border = f(), BORDER
        c.fill = CALC_FILL if kind == "calc" else INPUT_FILL
        if nf:
            c.number_format = nf
        if kind == "calc":
            cnt = lambda k: f'SUMIFS({rng(k)},{rng("case_study")},$B{r})'
            c.value = {
                "L": f'=IF($B{r}="","",COUNTIFS({rng("case_study")},$B{r}))',
                "M": f'=IF($B{r}="","",{cnt("reach_calc")})',
                "N": f'=IF($B{r}="","",{cnt("eng")})',
                "O": f'=IF(OR($B{r}="",N(M{r})=0),"",N{r}/M{r})',
                "P": f'=IF($B{r}="","",{cnt("clicks")})',
                "Q": f'=IF($B{r}="","",{cnt("sessions")})',
                "R": f'=IF($B{r}="","",{cnt("booked")})',
                "S": f'=IF($B{r}="","",{cnt("qualified")})',
                "T": f'=IF($B{r}="","",{cnt("pipeline")})',
            }[L]
    if lst:
        add_list_validation(ws_cs, f"{L}{CS_FIRST}:{L}{CS_LAST}", f"={LIST_REF[lst]}")
ws_cs.row_dimensions[4].height = 45
ws_cs.freeze_panes = "C5"
cs_example = ["CS01", "EXAMPLE - MRI no-show reduction", "cs01-mri-noshows", "Multi-site Imaging Network", "TX",
              "MRI", "EXAMPLE - high MRI no-show rate leaving slots empty", "EXAMPLE - automated reminders + waitlist backfill",
              "[replace with the real result, e.g. no-shows cut from X% to Y%]", "Hard Metric", ""]
for i, v in enumerate(cs_example, start=1):
    c = ws_cs.cell(CS_FIRST, i, v)
    c.font = f(italic=True, color="0000FF")
ws_cs.conditional_formatting.add(f"O{CS_FIRST}:O{CS_LAST}", DataBarRule(start_type="num", start_value=0, end_type="max", color="63BE7B"))
ws_cs.conditional_formatting.add(f"R{CS_FIRST}:R{CS_LAST}", DataBarRule(start_type="num", start_value=0, end_type="max", color="F4B183"))


# ------------------------------------------------------------ Follower Growth
title(ws_fg, "Follower Growth — weekly snapshot",
      "Enter the start Monday in A5 (other weeks fill themselves), then each Monday type the follower counts.")
FG_ACC = ["LinkedIn - Company", "LinkedIn - Founder", "Instagram", "X", "Facebook", "YouTube"]
ws_fg.cell(4, 1, "Week Of")
for i, a in enumerate(FG_ACC):
    ws_fg.cell(4, 2 + i, f"{a} Followers")
    ws_fg.cell(4, 2 + len(FG_ACC) + i, f"{a} Net New")
ws_fg.cell(4, 2 + 2 * len(FG_ACC), "Total Followers")
ws_fg.cell(4, 3 + 2 * len(FG_ACC), "Total Net New")
last_fg_col = 3 + 2 * len(FG_ACC)
for c in range(1, last_fg_col + 1):
    h = ws_fg.cell(4, c)
    is_calc = c == 1 and False or c > 1 + len(FG_ACC)
    h.font, h.fill = f(True, "FFFFFF"), fill(GREEN if is_calc else "2F5597")
    h.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws_fg.column_dimensions[get_column_letter(c)].width = 13
ws_fg.row_dimensions[4].height = 45
for r in range(5, 5 + 104):
    a = ws_fg.cell(r, 1)
    a.number_format, a.font, a.border = FMT_DATE, f(), BORDER
    if r == 5:
        a.value, a.fill = dt.date(2026, 9, 28), INPUT_FILL
        a.font = f(color="0000FF")
    else:
        a.value, a.fill = f"=A{r - 1}+7", CALC_FILL
    for i in range(len(FG_ACC)):
        inp = ws_fg.cell(r, 2 + i)
        inp.fill, inp.number_format, inp.font, inp.border = INPUT_FILL, FMT_INT, f(), BORDER
        L = get_column_letter(2 + i)
        net = ws_fg.cell(r, 2 + len(FG_ACC) + i)
        net.value = "" if r == 5 else f'=IF(OR({L}{r}="",{L}{r - 1}=""),"",{L}{r}-{L}{r - 1})'
        net.fill, net.number_format, net.font, net.border = CALC_FILL, "+#,##0;-#,##0;0", f(), BORDER
    first, lastc = get_column_letter(2), get_column_letter(1 + len(FG_ACC))
    t = ws_fg.cell(r, 2 + 2 * len(FG_ACC), f'=IF(COUNT({first}{r}:{lastc}{r})=0,"",SUM({first}{r}:{lastc}{r}))')
    nf1, nf2 = get_column_letter(2 + len(FG_ACC)), get_column_letter(1 + 2 * len(FG_ACC))
    tn = ws_fg.cell(r, 3 + 2 * len(FG_ACC), f'=IF(COUNT({nf1}{r}:{nf2}{r})=0,"",SUM({nf1}{r}:{nf2}{r}))')
    for cc, nf in ((t, FMT_INT), (tn, "+#,##0;-#,##0;0")):
        cc.fill, cc.number_format, cc.font, cc.border = CALC_FILL, nf, f(), BORDER
ws_fg.freeze_panes = "B5"


# ------------------------------------------------------------ Dashboard
title(ws_db, "Alpha Nodus — Social → Website → Demo Dashboard",
      "Everything below recalculates for the reporting period. Change the two yellow dates to switch periods.")
ws_db["B4"], ws_db["E4"] = "Period from:", "to:"
ws_db["C4"], ws_db["F4"] = dt.date(2026, 1, 1), dt.date(2026, 12, 31)
for a in ("C4", "F4"):
    ws_db[a].fill, ws_db[a].number_format, ws_db[a].font, ws_db[a].border = INPUT_FILL, FMT_DATE, f(True, "0000FF"), BORDER
for a in ("B4", "E4"):
    ws_db[a].font, ws_db[a].alignment = f(True), Alignment(horizontal="right")
FROM, TO = "$C$4", "$F$4"
DATE_CRIT = f'{rng("date")},">="&{FROM},{rng("date")},"<="&{TO}'
dp = "'Demo Pipeline'!"
DP_DATE_CRIT = f'{dp}$B${DP_FIRST}:$B${DP_LAST},">="&{FROM},{dp}$B${DP_FIRST}:$B${DP_LAST},"<="&{TO}'
ws_db.column_dimensions["A"].width = 2
ws_db.column_dimensions["B"].width = 34
for c in range(3, 17):
    ws_db.column_dimensions[get_column_letter(c)].width = 13


def section(r, text, color=NAVY):
    c = ws_db.cell(r, 2, text)
    c.font = f(True, color, 12)
    return r + 1


def tot(k):
    return f"SUMIFS({rng(k)},{DATE_CRIT})"


# KPI tiles
r = section(6, "Headline numbers for the period")
tiles = [
    ("Posts Published", f"=COUNTIFS({DATE_CRIT})", FMT_INT),
    ("Impressions", f"={tot('impr')}", FMT_INT),
    ("Reach (Calc.)", f"={tot('reach_calc')}", FMT_INT),
    ("Engagements", f"={tot('eng')}", FMT_INT),
    ("Engagement Rate", f'=IFERROR({tot("eng")}/{tot("reach_calc")},"")', FMT_PCT),
    ("Link Clicks", f"={tot('clicks')}", FMT_INT),
    ("Website Sessions", f"={tot('sessions')}", FMT_INT),
    ("Demo Page Visits", f"={tot('demo_page')}", FMT_INT),
    ("Demos Booked (from posts)", f"={tot('booked')}", FMT_INT),
    ("Demos Booked (ALL sources)", f"=COUNTIFS({DP_DATE_CRIT})", FMT_INT),
    ("% of Demos Traced to a Post", f'=IFERROR(SUMPRODUCT(({dp}$B${DP_FIRST}:$B${DP_LAST}>={FROM})*({dp}$B${DP_FIRST}:$B${DP_LAST}<={TO})*(LEN({dp}$Q${DP_FIRST}:$Q${DP_LAST})>0))/COUNTIFS({DP_DATE_CRIT}),"")', FMT_PCT),
    ("Qualified Demos (from posts)", f"={tot('qualified')}", FMT_INT),
    ("Pipeline from Posts ($)", f"={tot('pipeline')}", FMT_USD),
    ("Ad Spend ($)", f"={tot('spend')}", FMT_USD),
]
for i, (lab, form, nf) in enumerate(tiles):
    row = r + (i // 7) * 3
    colx = 2 + (i % 7) * 2 if False else 3 + (i % 7) * 2 - 1
    lc = ws_db.cell(row, colx, lab)
    lc.font, lc.fill = f(True, "FFFFFF", 9), fill(TEAL)
    lc.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws_db.merge_cells(start_row=row, start_column=colx, end_row=row, end_column=colx + 1)
    vc = ws_db.cell(row + 1, colx, form)
    vc.font, vc.number_format = f(True, NAVY, 14), nf
    vc.alignment = Alignment(horizontal="center")
    vc.fill = fill("DEEAF6")
    ws_db.merge_cells(start_row=row + 1, start_column=colx, end_row=row + 1, end_column=colx + 1)
    ws_db.row_dimensions[row].height = 28
r += 6

# Funnel
r = section(r + 1, "Funnel: post → website → demo  (period)")
funnel = [
    ("Impressions", tot("impr")), ("Link Clicks", tot("clicks")), ("Website Sessions", tot("sessions")),
    ("Demo Page Visits", tot("demo_page")), ("Demo Form Starts", tot("form_starts")),
    ("Demos Booked", tot("booked")), ("Demos Attended", tot("attended")),
    ("Qualified Demos", tot("qualified")), ("Deals Won", tot("won")),
]
for j, h in enumerate(["Stage", "Count", "% of Previous Stage", "% of Impressions"]):
    c = ws_db.cell(r, 2 + j, h)
    c.font, c.fill, c.alignment = f(True, "FFFFFF"), fill(NAVY), Alignment(wrap_text=True, horizontal="center")
first_funnel = r + 1
for i, (lab, form) in enumerate(funnel):
    rr = r + 1 + i
    ws_db.cell(rr, 2, lab).font = f(True)
    ws_db.cell(rr, 3, f"={form}").number_format = FMT_INT
    ws_db.cell(rr, 4, "" if i == 0 else f'=IFERROR(C{rr}/C{rr - 1},"")').number_format = "0.0%"
    ws_db.cell(rr, 5, f'=IFERROR(C{rr}/C${first_funnel},"")').number_format = "0.00%"
    for cc in range(2, 6):
        ws_db.cell(rr, cc).border = BORDER
        if cc > 2:
            ws_db.cell(rr, cc).font = f()
ws_db.cell(r + 1 + len(funnel), 2, "Form Starts only fills if you track them in GA4; a 0 there breaks the next % — that's expected.").font = f(italic=True, color=GREY, size=9)
ws_db.conditional_formatting.add(f"C{first_funnel}:C{first_funnel + len(funnel) - 1}", DataBarRule(start_type="num", start_value=0, end_type="max", color="5B9BD5"))
r += len(funnel) + 3

# Top 5 tables
TOP_HDR = ["Rank", "Post ID", "Date", "Platform", "Account", "Content Pillar", "Case Study", "Format", "Eng. Rate", "Reach", "Link Clicks", "Demos Booked"]
TOP_KEYS = [None, "post_id", "date", "platform", "account", "pillar", "case_study", "format", "er", "reach_calc", "clicks", "booked"]
TOP_FMT = [None, None, FMT_DATE, None, None, None, None, None, FMT_PCT, FMT_INT, FMT_INT, FMT_INT]
HELPER_COL = 16   # column P holds the LARGE() key, hidden
for label, rkey in (("Top 5 posts by Engagement Rate (period)", "rk_er"), ("Top 5 posts by Demos Booked (period)", "rk_demo")):
    r = section(r, label, GREEN)
    for j, h in enumerate(TOP_HDR):
        c = ws_db.cell(r, 2 + j, h)
        c.font, c.fill, c.alignment = f(True, "FFFFFF"), fill(GREEN), Alignment(wrap_text=True, horizontal="center")
    for n in range(1, 6):
        rr = r + n
        key = ws_db.cell(rr, HELPER_COL, f'=IFERROR(LARGE({rng(rkey)},{n}),"")')
        key.font = f(color=GREY, size=8)
        kref = f"${get_column_letter(HELPER_COL)}{rr}"
        for j, k in enumerate(TOP_KEYS):
            c = ws_db.cell(rr, 2 + j)
            c.border, c.font = BORDER, f()
            if k is None:
                c.value = n
            else:
                c.value = f'=IF({kref}="","",INDEX({rng(k)},MATCH({kref},{rng(rkey)},0))&"")' if TOP_FMT[j] is None else \
                          f'=IF({kref}="","",INDEX({rng(k)},MATCH({kref},{rng(rkey)},0)))'
                if TOP_FMT[j]:
                    c.number_format = TOP_FMT[j]
    r += 7
ws_db.column_dimensions[get_column_letter(HELPER_COL)].hidden = True

# Breakdown tables
DIM_HDR = ["Posts", "Impressions", "Reach", "Engagements", "Eng. Rate", "Link Clicks", "CTR", "Sessions",
           "Demo Page Visits", "Demos Booked", "Qualified Demos", "Session → Booked %", "Pipeline ($)"]


def dim_table(r, label, dim_key, labels, note=None):
    r = section(r, label)
    if note:
        ws_db.cell(r - 1, 5, note).font = f(italic=True, color=GREY, size=9)
    ws_db.cell(r, 2, "").fill = fill(NAVY)
    for j, h in enumerate(DIM_HDR):
        c = ws_db.cell(r, 3 + j, h)
        c.font, c.fill, c.alignment = f(True, "FFFFFF"), fill(NAVY), Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws_db.row_dimensions[r].height = 30
    first = r + 1
    for i, lab in enumerate(labels):
        rr = first + i
        lc = ws_db.cell(rr, 2, lab)
        lc.font, lc.border = f(True), BORDER
        crit = f'{rng(dim_key)},$B{rr},{DATE_CRIT}'
        s = lambda k: f"SUMIFS({rng(k)},{crit})"
        g = lambda body: f'=IF($B{rr}="","",{body})'
        forms = [
            g(f"COUNTIFS({crit})"), g(s("impr")), g(s("reach_calc")), g(s("eng")),
            g(f'IFERROR(F{rr}/E{rr},"")'), g(s("clicks")), g(f'IFERROR(H{rr}/D{rr},"")'), g(s("sessions")),
            g(s("demo_page")), g(s("booked")), g(s("qualified")), g(f'IFERROR(L{rr}/J{rr},"")'), g(s("pipeline")),
        ]
        fmts = [FMT_INT, FMT_INT, FMT_INT, FMT_INT, FMT_PCT, FMT_INT, FMT_PCT, FMT_INT, FMT_INT, FMT_INT, FMT_INT, FMT_PCT, FMT_USD]
        for j, (form, nf) in enumerate(zip(forms, fmts)):
            c = ws_db.cell(rr, 3 + j, form)
            c.number_format, c.border, c.font = nf, BORDER, f()
    last = first + len(labels) - 1
    ws_db.conditional_formatting.add(f"G{first}:G{last}", DataBarRule(start_type="num", start_value=0, end_type="max", color="63BE7B"))
    ws_db.conditional_formatting.add(f"L{first}:L{last}", DataBarRule(start_type="num", start_value=0, end_type="max", color="F4B183"))
    return last + 2


def labels_of(name, spare=2):
    L = LIST_COL[name]
    return [f'=IF(Settings!${L}${5 + i}="","",Settings!${L}${5 + i})' for i in range(len(LISTS[name]) + spare)]


r = dim_table(r, "By Platform", "platform", labels_of("Platform"))
r = dim_table(r, "By Account (Company page vs Founder vs Employees)", "account", labels_of("Account"))
r = dim_table(r, "By Content Pillar", "pillar", labels_of("Content Pillar"))
cs_labels = [f"=IF('Case Studies'!$B${CS_FIRST + i}=\"\",\"\",'Case Studies'!$B${CS_FIRST + i})" for i in range(15)]
r = dim_table(r, "By Case Study (first 15 on the Case Studies sheet)", "case_study", cs_labels,
              "Full all-time list with totals lives on the Case Studies sheet.")
r = dim_table(r, "By Post Format", "format", labels_of("Post Format"))
r = dim_table(r, "By Target Persona", "persona", labels_of("Target Persona"))
r = dim_table(r, "By Hook Type", "hook", labels_of("Hook Type"))
r = dim_table(r, "By CTA", "cta", labels_of("CTA"))
r = dim_table(r, "By Link Placement", "placement", labels_of("Link Placement"))
r = dim_table(r, "By Day of Week", "day", LISTS["Day"])
r = dim_table(r, "Organic vs Paid", "paid", labels_of("Paid Status"))

# Monthly trend (12 months from the period start)
r = section(r, "Monthly trend (12 months starting from the 'Period from' month)")
ws_db.cell(r, 2, "Month").font = f(True, "FFFFFF")
ws_db.cell(r, 2).fill = fill(NAVY)
TREND = [("Posts", None, FMT_INT), ("Reach", "reach_calc", FMT_INT), ("Engagements", "eng", FMT_INT),
         ("Eng. Rate", "__er", FMT_PCT), ("Link Clicks", "clicks", FMT_INT), ("Sessions", "sessions", FMT_INT),
         ("Demo Page Visits", "demo_page", FMT_INT), ("Demos Booked (posts)", "booked", FMT_INT),
         ("Demos Booked (all sources)", "__dpall", FMT_INT), ("New Followers", "new_followers", FMT_INT)]
for j, (h, _, _) in enumerate(TREND):
    c = ws_db.cell(r, 3 + j, h)
    c.font, c.fill, c.alignment = f(True, "FFFFFF"), fill(NAVY), Alignment(wrap_text=True, horizontal="center")
ws_db.row_dimensions[r].height = 30
for i in range(12):
    rr = r + 1 + i
    mc = ws_db.cell(rr, 2, f'=TEXT(DATE(YEAR({FROM}),MONTH({FROM})+{i},1),"yyyy-mm")')
    mc.font, mc.border = f(True), BORDER
    for j, (h, k, nf) in enumerate(TREND):
        if k is None:
            form = f'=COUNTIFS({rng("month")},$B{rr})'
        elif k == "__er":
            form = f'=IFERROR(E{rr}/D{rr},"")'
        elif k == "__dpall":
            form = (f'=COUNTIFS({dp}$B${DP_FIRST}:$B${DP_LAST},">="&DATE(YEAR({FROM}),MONTH({FROM})+{i},1),'
                    f'{dp}$B${DP_FIRST}:$B${DP_LAST},"<"&DATE(YEAR({FROM}),MONTH({FROM})+{i + 1},1))')
        else:
            form = f'=SUMIFS({rng(k)},{rng("month")},$B{rr})'
        c = ws_db.cell(rr, 3 + j, form)
        c.number_format, c.border, c.font = nf, BORDER, f()
ws_db.conditional_formatting.add(f"J{r + 1}:J{r + 12}", DataBarRule(start_type="num", start_value=0, end_type="max", color="F4B183"))
r += 15

# Demo pipeline by source & attribution
r = section(r, "Demos by Lead Source — all channels (period, from Demo Pipeline)", ORANGE)
SRC_HDR = ["Demos Booked", "Attended", "Qualified", "Closed Won", "Pipeline ($)", "Avg Days Post → Booking"]
for j, h in enumerate(SRC_HDR):
    c = ws_db.cell(r, 3 + j, h)
    c.font, c.fill, c.alignment = f(True, "FFFFFF"), fill(ORANGE), Alignment(wrap_text=True, horizontal="center")
ws_db.cell(r, 2).fill = fill(ORANGE)
ws_db.row_dimensions[r].height = 30


def dp_rows(r, labels, colL):
    for i, lab in enumerate(labels):
        rr = r + 1 + i
        ws_db.cell(rr, 2, lab).font = f(True)
        crit = f'{dp}${colL}${DP_FIRST}:${colL}${DP_LAST},$B{rr},{DP_DATE_CRIT}'
        forms = [f"=COUNTIFS({crit})",
                 f'=COUNTIFS({crit},{dp}$W${DP_FIRST}:$W${DP_LAST},"Attended")',
                 f'=COUNTIFS({crit},{dp}$X${DP_FIRST}:$X${DP_LAST},"Yes")',
                 f'=COUNTIFS({crit},{dp}$Y${DP_FIRST}:$Y${DP_LAST},"Closed Won")',
                 f"=SUMIFS({dp}$Z${DP_FIRST}:$Z${DP_LAST},{crit})",
                 f'=IFERROR(AVERAGEIFS({dp}$U${DP_FIRST}:$U${DP_LAST},{crit}),"")']
        for j, (form, nf) in enumerate(zip(forms, [FMT_INT] * 4 + [FMT_USD, "0.0"])):
            c = ws_db.cell(rr, 3 + j, form)
            c.number_format, c.border, c.font = nf, BORDER, f()
        ws_db.cell(rr, 2).border = BORDER
    return r + len(labels) + 2


r = dp_rows(r, labels_of("Lead Source"), "J")
r = section(r, "Demos by Attribution Type — how well is tracking working?", ORANGE)
for j, h in enumerate(SRC_HDR):
    c = ws_db.cell(r, 3 + j, h)
    c.font, c.fill, c.alignment = f(True, "FFFFFF"), fill(ORANGE), Alignment(wrap_text=True, horizontal="center")
ws_db.cell(r, 2).fill = fill(ORANGE)
r = dp_rows(r, ["UTM tracked", "Manual / self-reported", "Self-reported channel only", "Unattributed"], "R")
ws_db.cell(r - 1, 2, "If 'Unattributed' is large, fix the hidden UTM fields on the demo form and make 'How did you hear about us?' required.").font = f(italic=True, color=GREY, size=9)
ws_db.freeze_panes = "A5"


# ------------------------------------------------------------ Start Here
title(ws_start, "Alpha Nodus — Social Media Performance & Demo Attribution Tracker",
      "Built for LinkedIn (company + founder), Instagram, X (plus Facebook / YouTube if you use them). Audience: independent imaging & radiology centers, USA.")
ws_start.column_dimensions["A"].width = 3
ws_start.column_dimensions["B"].width = 30
ws_start.column_dimensions["C"].width = 110
rows = [
    ("h", "What this workbook answers"),
    ("", "1. Which platform, content pillar, case study, format, hook and persona gets the best reach and engagement?"),
    ("", "2. Which posts actually send imaging-center decision-makers to alphanodus.com, to the demo page, and into a booked demo?"),
    ("", "3. What does each demo cost (if boosted) and how much pipeline did each post / case study create?"),
    ("h", "Colour legend"),
    ("legend_in", "Yellow cell = you type here (or pick from the dropdown)."),
    ("legend_calc", "Grey cell = formula. Don't type over it — it recalculates by itself."),
    ("legend_ex", "Blue italic text = EXAMPLE values in row 5 of Post Tracker, Demo Pipeline and Case Studies. Overwrite or delete them before real use."),
    ("", "Hover over any column header with a red corner triangle for a note on what to enter."),
    ("h", "Sheet map"),
    ("Post Tracker", "One row per post per platform. Details → auto tracking link → platform metrics → auto reach & engagement → website/demo funnel."),
    ("Demo Pipeline", "One row per demo booked from ANY source. Its UTM / manual Post ID links the demo back to the post automatically."),
    ("Dashboard", "Pick a date range. Headline KPIs, funnel, top 5 posts, and breakdowns by platform, account, pillar, case study, format, persona, hook, CTA, link placement, weekday, paid vs organic, monthly trend and demo sources."),
    ("Case Studies", "Library of case studies (name, slug, client type, problem, result). Shows all-time reach, engagement and demos for each story."),
    ("Follower Growth", "Weekly follower count per account → net new followers."),
    ("Tracking Setup", "UTM rules, where each number comes from (GA4 paths), website set-up checklist, and every attribution scenario."),
    ("Metric Definitions", "Exact formula for every calculated metric and platform-by-platform notes."),
    ("Settings", "Dropdown lists, default landing page and engagement weights. Add new pillars / formats / personas here."),
    ("h", "Workflow"),
    ("Before posting", "Add a row: date, time, platform, account, format, pillar, case study, topic, hook, persona, CTA. The Post ID and Tracking URL appear automatically."),
    ("", "Copy the Tracking URL (or shorten it and paste the short link in the Short Link column) and use it as the ONLY link in the post / first comment / bio / story sticker."),
    ("Right after posting", "Paste the Live Post URL and the account's current follower count."),
    ("Day 7 after posting", "Open the platform's post analytics and fill the purple Platform Metrics block. Always pull at the same age so posts compare fairly. Count ICP engagers."),
    ("Weekly (Mondays)", "GA4 → fill Sessions, Engaged Sessions, Demo Page Visits for each post (see Tracking Setup). Update Follower Growth."),
    ("Every demo booked", "Add a row to Demo Pipeline with the utm_* values from the form / CRM and the 'How did you hear about us?' answer. Update status, qualified and stage as it progresses."),
    ("Monthly review", "Dashboard: which pillar / case study / format has the highest engagement rate AND demos? Write 'What We Learned' + 'Next Action' on the best and worst posts. Scale winners, retire losers."),
    ("h", "Rules that keep the data clean"),
    ("", "• Every post gets its own tracking link — never reuse a link across posts or platforms (the Post ID is the link's fingerprint)."),
    ("", "• Post ID = platform code + account code + date + time (e.g. LIC-260922-0900). A red Post ID means a duplicate — change the time by a minute."),
    ("", "• Never put UTM links on alphanodus.com's own internal links — it resets the visitor's source."),
    ("", "• One post published on 3 platforms = 3 rows (each gets its own link), sharing the same case study so they roll up together."),
    ("", "• Judge a post on engagement rate (not raw likes) and on demos; judge a case study only after it has run on at least 2–3 posts."),
]
r = 4
for kind, text in rows:
    if kind == "h":
        r += 1
        c = ws_start.cell(r, 2, text)
        c.font = f(True, NAVY, 12)
    elif kind.startswith("legend"):
        sw = ws_start.cell(r, 2, {"legend_in": "Input", "legend_calc": "Calculated", "legend_ex": "Example"}[kind])
        sw.fill = {"legend_in": INPUT_FILL, "legend_calc": CALC_FILL, "legend_ex": INPUT_FILL}[kind]
        sw.font = f(True, "0000FF", italic=True) if kind == "legend_ex" else f(True)
        sw.border = BORDER
        ws_start.cell(r, 3, text).font = f()
    else:
        ws_start.cell(r, 2, kind).font = f(True)
        c = ws_start.cell(r, 3, text)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1


# ------------------------------------------------------------ Tracking Setup
title(ws_ut, "Tracking Setup — UTMs, data sources, website checklist, attribution scenarios",
      "Work through the checklist once with whoever manages alphanodus.com; the rest is reference.")
widths = [3, 28, 44, 44, 40, 14, 16]
for i, w in enumerate(widths, start=1):
    ws_ut.column_dimensions[get_column_letter(i)].width = w


def table(r, heading, headers, data, color=NAVY, status_col=None):
    ws_ut.cell(r, 2, heading).font = f(True, color, 12)
    r += 1
    for j, h in enumerate(headers):
        c = ws_ut.cell(r, 2 + j, h)
        c.font, c.fill, c.alignment = f(True, "FFFFFF"), fill(color), Alignment(wrap_text=True, vertical="center")
    for row in data:
        r += 1
        for j, v in enumerate(row):
            c = ws_ut.cell(r, 2 + j, v)
            c.font, c.border = f(bold=(j == 0)), BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if status_col is not None:
            for sc in status_col:
                cc = ws_ut.cell(r, 2 + sc)
                cc.fill = INPUT_FILL
    return r + 2


r = 4
r = table(r, "1. UTM naming rules (the tracker builds these for you)", ["Parameter", "Rule", "Example", "Why"], [
    ("utm_source", "Platform, lower-case: linkedin · instagram · twitter · facebook · youtube", "linkedin",
     "Uses 'twitter' for X because GA4's built-in social-source list recognises it, so X traffic lands in the Social channel group."),
    ("utm_medium", "social (organic) or paid_social (boosted / paid)", "social",
     "These two values map cleanly onto GA4's 'Organic Social' and 'Paid Social' channel groups."),
    ("utm_campaign", "Case study slug if the post features one, otherwise the content-pillar slug", "cs01-mri-noshows",
     "Lets GA4 roll up every post about the same story across all platforms."),
    ("utm_content", "The Post ID, lower-case", "lic-260922-0900",
     "Unique per post — this is the key that ties a website visit and a demo back to one exact post."),
    ("Short links", "Optional. Shorten the full Tracking URL with bit.ly / Dub / a branded domain (e.g. go.alphanodus.com)",
     "go.alphanodus.com/mri-cs", "Cleaner in posts and gives a second click count. The UTM must still be inside the destination URL."),
])
r = table(r, "2. Where each number comes from", ["Tracker column(s)", "Source", "How to pull it", "Notes"], [
    ("Impressions, Reach, Reactions, Comments, Reposts, Clicks", "LinkedIn",
     "Company page: Analytics → Content → the post's metrics (or 'View analytics' under the post). Founder profile: the analytics link under the post.",
     "Company-page posts show 'Members reached'; personal profiles also show impressions and members reached."),
    ("Views, Reach, Likes, Comments, Shares, Saves, Profile visits, Follows", "Instagram",
     "Professional dashboard / 'View insights' under the post.", "Instagram now reports 'Views' as its headline number — enter it in Impressions / Views. Captions aren't clickable: track link-in-bio and story-sticker clicks."),
    ("Impressions, Likes, Replies, Reposts, Bookmarks, Link clicks, Profile visits", "X",
     "Click the analytics (bar-chart) icon on the post.", "X shows no unique reach — the tracker uses impressions as the reach proxy and labels it."),
    ("Website Sessions, Engaged Sessions", "Google Analytics 4",
     "Reports → Acquisition → Traffic acquisition → add a secondary dimension 'Session manual ad content' → search the Post ID (lower-case). Or Explore → Free form: rows = Session manual ad content, values = Sessions, Engaged sessions.",
     "Expect sessions < link clicks (bots, link previews, people bouncing before the page loads). Click → Session % shows the gap."),
    ("Demo Page Visits", "Google Analytics 4",
     "Explore → Free form: rows = Session manual ad content; values = Sessions; filter Page path contains your demo page path (e.g. /demo or /book-a-demo). Better: set up the 'demo_page_view' key event (checklist) and read the key-event count.",
     None),
    ("Demos Booked / Attended / Qualified / Pipeline", "Demo form + CRM → Demo Pipeline sheet",
     "Each booking's hidden UTM fields (or the CRM's original source fields) are pasted into Demo Pipeline. The Post Tracker counts them automatically.",
     "GA4 conversions are for trend-checking; the Demo Pipeline sheet is the source of truth for demos."),
    ("Followers", "Each platform", "Follower count shown on the profile / page analytics every Monday.", None),
])
r = table(r, "3. Website set-up checklist (one-time)", ["Step", "What to do", "Why it matters", "Tip", "Status", "Owner"], [
    ("GA4 on every page", "Confirm the GA4 tag fires on all alphanodus.com pages, including the demo page and the booking confirmation page.", "No tag = no sessions to attribute.", "Check with GA4 → Admin → DebugView or Google Tag Assistant."),
    ("Demo page key event", "Create a GA4 event 'demo_page_view' when page_location contains the demo page path, and mark it as a key event.", "Gives the Demo Page Visits number directly per utm_content.", "GA4 → Admin → Events → Create event."),
    ("Demo booked key event", "Fire 'demo_booked' on the booking confirmation / thank-you page (or on the scheduler's 'event scheduled' message if the calendar is embedded) and mark it as a key event.", "Trend-checks bookings inside GA4 by source / campaign / content.", "Calendly embeds post a 'calendly.event_scheduled' browser message that Google Tag Manager can listen for; HubSpot meetings have a similar callback."),
    ("Hidden UTM fields on the demo form", "Add hidden fields utm_source, utm_medium, utm_campaign, utm_content (+ utm_term) to the demo form and fill them from the URL.", "This is what lets a demo row be traced to an exact post.", "Most form tools (HubSpot, Typeform, Jotform, Calendly) support UTM capture natively."),
    ("Persist UTMs across pages & visits", "Save the first UTM values in a first-party cookie / localStorage (30–90 days) and write them into the form even if the visitor browsed other pages or came back later.", "Visitors rarely book on the first page — without this, the UTM is lost by the time they reach the form.", "Keep both first-touch and last-touch values if your CRM allows."),
    ("'How did you hear about us?'", "Add a required dropdown / free-text field to the demo form.", "Catches 'dark social': people who saw a post but typed the URL or searched Google.", "Options match the 'How Heard' list on Settings."),
    ("CRM source fields", "Map the hidden UTM fields into CRM contact / deal properties (original source, campaign, content).", "Lets you follow demo → qualified → closed won in the Demo Pipeline sheet.", None),
    ("LinkedIn Insight Tag", "Install the LinkedIn Insight Tag on the site.", "Shows the job titles, industries and company sizes of site visitors, and is required for LinkedIn conversion tracking / retargeting.", "LinkedIn Campaign Manager → Analyze → Insight Tag."),
    ("Meta Pixel / X Pixel", "Install only if you plan to run Instagram / Facebook / X ads or retargeting.", "Paid-social conversion tracking and audiences.", None),
    ("Exclude internal traffic", "Define your team's IP addresses as internal traffic in GA4 and filter them out.", "Team clicks on your own posts inflate sessions.", "GA4 → Admin → Data streams → Configure tag settings → Define internal traffic."),
    ("Data retention", "Set GA4 event data retention to 14 months.", "Default is 2 months, which limits year-over-year Explorations.", "GA4 → Admin → Data collection and modification → Data retention."),
    ("Booking tool on another domain?", "If the scheduler opens on a different domain, embed it on alphanodus.com or configure GA4 cross-domain measurement.", "Otherwise the booking shows up as a new 'referral' session and the post loses credit.", None),
    ("Profile / bio links", "Give every profile 'website' link its own UTM (utm_content = profile-linkedin-company, profile-instagram, …).", "Separates 'clicked the profile link' from 'clicked a post link'.", None),
    ("Cookie consent", "If a consent banner is used, expect GA4 to under-count vs platform link clicks.", "Explains a lower Click → Session % — not a tracking bug.", None),
], status_col=[4, 5])
dv_status = DataValidation(type="list", formula1=f"={LIST_REF['Checklist Status']}", allow_blank=True)
ws_ut.add_data_validation(dv_status)
dv_status.add(f"F{r - 15}:F{r - 2}")
r = table(r, "4. Attribution scenarios — what happens, and how the tracker captures it", ["Scenario", "What the data shows", "How we capture it", "Where it lands"], [
    ("Clicks the post link and books a demo in the same visit", "GA4 session with utm_content = Post ID; form hidden fields filled.", "UTM copied into Demo Pipeline → auto-linked to the post.", "Demo Pipeline → Attribution Type 'UTM tracked'."),
    ("Clicks the post link, leaves, comes back days later via Google / typing the URL, then books", "GA4 last-click credits Google / Direct.", "Persisted first-touch UTM cookie fills the hidden fields; CRM original source; self-reported field.", "UTM tracked (if persisted) or Manual Post ID."),
    ("Sees the post but never clicks, later types alphanodus.com (dark social)", "Direct traffic; no UTM anywhere.", "'How did you hear about us?' answer; ask on the demo call which post they saw and add the Manual Post ID.", "Self-reported / Manual; watch for Direct-traffic bumps on posting days."),
    ("Engages with the post, then DMs or comments and the team books the demo manually", "Nothing in GA4.", "Log the demo with Lead Source 'Social DM' and the Post ID in Manual Post ID; log the DM in 'Inbound DMs'.", "Manual / self-reported."),
    ("Visits the company page / profile after the post, clicks the website link there", "UTM from the profile link, not the post.", "Profile links carry their own utm_content (profile-…).", "Shows as the profile link — correlate with the post's Profile Visits column."),
    ("Instagram: link in bio or story sticker", "UTM from the bio / sticker link.", "Use a per-post link in the bio link page or sticker (the post's Tracking URL); set Link Placement accordingly.", "UTM tracked."),
    ("Someone forwards / DMs the post link to a colleague who books", "UTM travels with the link.", "Credit goes to the original post — correct.", "UTM tracked."),
    ("Founder or employee reshares the company post", "Their audience is different.", "Log the reshare as its own row (Account = Founder / Employee) with its own tracking link.", "Separate row; compare by Account on the Dashboard."),
    ("Sees several posts over weeks before booking (multi-touch)", "Only one post gets the UTM credit.", "Keep first- and last-touch UTMs in the CRM; use the self-reported answer and 'Days: Post → Booking' to judge lag.", "Last-touch in Demo Pipeline; note others in Notes."),
    ("Boosted / paid version of a post", "utm_medium = paid_social.", "Mark Organic / Paid, enter Ad Spend → Cost per Demo calculates.", "Organic vs Paid table on the Dashboard."),
    ("Link clicks are high but sessions are low", "Bots, previews, slow page or consent-banner loss.", "Watch Click → Session %; check page speed on mobile.", "Post Tracker."),
    ("Books a demo from a non-social source (Google, referral, event)", "No social UTM.", "Still log it in Demo Pipeline so you can see what share of all demos social drives.", "Dashboard → '% of Demos Traced to a Post' and 'Demos by Lead Source'."),
], color=ORANGE)
r = table(r, "5. Monthly review questions", ["Question", "Where to look", "", ""], [
    ("Which platform gives the best engagement rate AND demos?", "Dashboard → By Platform", None, None),
    ("Do founder posts beat company-page posts?", "Dashboard → By Account", None, None),
    ("Which case study / pillar should we make more of?", "Dashboard → By Case Study, By Content Pillar; Case Studies sheet", None, None),
    ("Which format, hook and CTA work for imaging-center owners vs radiologists vs ops managers?", "Dashboard → By Format / Hook / CTA / Persona", None, None),
    ("Are we reaching buyers or just peers?", "Post Tracker → ICP Engagers; LinkedIn Insight Tag visitor demographics", None, None),
    ("Where does the funnel leak?", "Dashboard → Funnel (% of previous stage)", None, None),
    ("How long from post to demo?", "Demo Pipeline → Days: Post → Booking; Dashboard → Avg Days", None, None),
    ("Is tracking healthy?", "Dashboard → Demos by Attribution Type ('Unattributed' should shrink)", None, None),
], color=GREEN)


# ------------------------------------------------------------ Metric Definitions
title(ws_md, "Metric Definitions", "How every calculated number is worked out, so the whole team reads them the same way.")
ws_md.column_dimensions["A"].width = 3
ws_md.column_dimensions["B"].width = 28
ws_md.column_dimensions["C"].width = 55
ws_md.column_dimensions["D"].width = 70
defs = [
    ("Metric", "Formula", "What it tells you"),
    ("Reach (Calc.)", "Reported reach; if blank, Impressions / Views", "Unique people who saw the post. X doesn't report reach, so impressions stand in (flagged in Reach Basis)."),
    ("Reach % of Followers", "Reach (Calc.) ÷ Followers at Posting", "How much of your audience the algorithm showed it to; above 100% = it travelled beyond followers."),
    ("Total Engagements", "Likes + Comments + Shares + Saves + Link Clicks", "All active interactions. Views, profile visits and follows are tracked separately."),
    ("Engagement Rate (by Reach)", "Total Engagements ÷ Reach (Calc.)", "THE main quality metric — compares posts of different sizes fairly."),
    ("Engagement Rate (by Impr.)", "Total Engagements ÷ Impressions", "Matches how LinkedIn's own dashboard reports engagement rate."),
    ("Click-Through Rate", "Link Clicks ÷ Impressions", "How well the post drives traffic to alphanodus.com."),
    ("Share Rate", "Shares ÷ Reach (Calc.)", "Word-of-mouth potential — shared posts reach peers at other imaging centers."),
    ("Weighted Eng. per 1K Reach", "(Likes×w1 + Comments×w2 + Shares×w3 + Saves×w4 + Clicks×w5) ÷ Reach × 1000", "Values deeper actions more than likes. Weights are on Settings."),
    ("ER vs Platform Avg", "Post Engagement Rate ÷ average Engagement Rate of all posts on that platform", "1.2x or more = clear winner, 0.8x or less = under-performer. Fair because each platform has different norms."),
    ("ICP Engagers", "Manual count", "Engagers who are imaging / radiology decision-makers — quality over quantity."),
    ("Click → Session %", "GA4 Sessions ÷ Link Clicks", "How many clicks turned into a real website visit."),
    ("Session → Demo Page %", "Demo Page Visits ÷ Sessions", "How well the landing page moves visitors toward a demo."),
    ("Demo Page → Booked %", "Demos Booked ÷ Demo Page Visits", "How well the demo page / form converts."),
    ("Session → Booked %", "Demos Booked ÷ Sessions", "End-to-end website conversion for that post's traffic."),
    ("Cost per Demo", "Ad Spend ÷ Demos Booked", "Only for boosted / paid posts."),
    ("Demos Booked (per post)", "Count of Demo Pipeline rows whose Attributed Post ID = Post ID", "Attributed Post ID = utm_content if present, else the Manual Post ID."),
    ("% of Demos Traced to a Post", "Demos with an Attributed Post ID ÷ all demos in the period", "How much of the demo pipeline social can prove it created."),
]
for i, row in enumerate(defs):
    for j, v in enumerate(row):
        c = ws_md.cell(4 + i, 2 + j, v)
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 0:
            c.font, c.fill = f(True, "FFFFFF"), fill(NAVY)
        else:
            c.font = f(bold=(j == 0))

# Tab colours & sheet order already set
for ws, color in ((ws_start, NAVY), (ws_pt, "2F5597"), (ws_dp, ORANGE), (ws_db, GREEN), (ws_cs, TEAL),
                  (ws_fg, PURPLE), (ws_ut, GREY), (ws_md, GREY), (ws_set, "A6A6A6")):
    ws.sheet_properties.tabColor = color
    ws.sheet_view.zoomScale = 90
ws_start.sheet_view.showGridLines = False
ws_db.sheet_view.showGridLines = False
ws_ut.sheet_view.showGridLines = False
ws_md.sheet_view.showGridLines = False

wb.save(OUT)
print(f"saved {OUT}  (Post Tracker columns: {PT_LASTCOL})")
