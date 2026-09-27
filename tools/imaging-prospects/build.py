"""Merge shortlist + activity check -> pick 10 Voices + 10 Operators per role -> workbook + profile dossiers."""
import csv, json, math, re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DATA = Path(sys.argv[1])  # work dir: raw/, activity.tsv, shortlist.json, overrides.json
OUT = Path(sys.argv[2])
(OUT / "profiles").mkdir(parents=True, exist_ok=True)

ROLE_LABEL = {"rcm": "Revenue Cycle Manager",
              "referral": "Referral Acquisition & Growth Lead",
              "growth": "Practice Growth Manager"}
PER_SEGMENT = 10

# People the web check showed have left the role / are not a fit (see activity.tsv evidence)
OV = json.loads((DATA / "overrides.json").read_text())
DROP = set(OV["drop"])                                  # left the role / not a fit (see activity.tsv evidence)
DEDUPE = {tuple(x) for x in OV["dedupe"]}               # (role, name) pairs removed when a person appears in two pulls
BONUS = OV["bonus"]                                     # evidence-strength bonus: podcast host, speaker, advocacy, recency
FLAG = OV["flag"]                                       # "verify" notes for newer web headlines

activity = {r["name"]: r for r in csv.DictReader(open(DATA / "activity.tsv"), delimiter="\t")}
shortlist = json.loads((DATA / "shortlist.json").read_text())
raw = {}
for f in ("rcm", "referral", "growth"):
    for p in json.loads((DATA / "raw" / f"{f}.json").read_text())["profiles"]:
        raw[p["person_id"]] = p


def voice_score(r):
    a = activity.get(r["name"], {})
    lvl = {"A": 10, "B": 5}.get(a.get("activity_level"), 0)
    reach = r["followers"] or r["connections"] or 0
    return round(lvl + BONUS.get(r["name"], 0) + 1.5 * math.log10(reach + 10), 2)


for rows in shortlist.values():
    for r in rows:
        r["followers"] = r["followers"] or None

selected = {}
for role, rows in shortlist.items():
    rows = [r for r in rows if r["name"] not in DROP and (role, r["name"]) not in DEDUPE]
    for r in rows:
        a = activity.get(r["name"], {})
        r["activity_level"] = a.get("activity_level", "not checked")
        r["evidence"] = a.get("evidence", "")
        r["vanity_url"] = a.get("vanity_url", "")
        r["voice_score"] = voice_score(r)
        r["flag"] = FLAG.get(r["name"], "")
    voices = sorted([r for r in rows if r["activity_level"] in ("A", "B")], key=lambda r: -r["voice_score"])[:PER_SEGMENT]
    vnames = {r["name"] for r in voices}
    operators = sorted([r for r in rows if r["name"] not in vnames], key=lambda r: -r["work_score"])[:PER_SEGMENT]
    for r in voices: r["segment"] = "Voice (active on LinkedIn)"
    for r in operators: r["segment"] = "Operator (top by track record)"
    selected[role] = voices + operators
    print(role, "voices", len(voices), "operators", len(operators), file=sys.stderr)


# ---------- profile dossiers ----------
def dt(s):
    return (s or "")[:7] or "?"


def dossier(r, rank):
    p = raw[r["person_id"]]
    L = [f"# {p['name']}", "", f"**{p.get('headline') or ''}**", "",
         f"- Role bucket: {ROLE_LABEL[r['role']]} | Segment: {r['segment']} | Rank {rank}",
         f"- Location: {p.get('region')}",
         f"- LinkedIn: {r['vanity_url'] or p['linkedin_profile_url']}",
         f"- LinkedIn (ID link, opens when logged in): {p['linkedin_profile_url']}",
         f"- Followers: {p.get('num_of_followers') or 'n/a'} | Connections: {p.get('num_of_connections')} | "
         f"Years of experience: {p.get('years_of_experience_raw')}",
         f"- LinkedIn activity: **{r['activity_level']}** - {r['evidence']}"]
    if r["flag"]:
        L.append(f"- Verify: {r['flag']}")
    L += ["", "## Current roles"]
    for e in p.get("current_employers") or []:
        L.append(f"- **{e.get('title')}**, {e.get('name')} (since {dt(e.get('start_date'))}; "
                 f"{e.get('years_at_company_raw', '?')} yrs; company size {e.get('company_headcount_latest', '?')}; "
                 f"{e.get('company_website_domain') or ''})")
    past = p.get("past_employers") or []
    if past:
        L += ["", "## Past roles"]
        for e in past[:12]:
            L.append(f"- {e.get('title')}, {e.get('name')} ({dt(e.get('start_date'))} to {dt(e.get('end_date'))})")
    edu = [e for e in p.get("education_background") or [] if e.get("institute_name")]
    if edu:
        L += ["", "## Education"]
        for e in edu:
            L.append(f"- {e.get('institute_name')}: {' - '.join(x for x in (e.get('degree_name'), e.get('field_of_study')) if x)}")
    if p.get("skills"):
        L += ["", "## Skills", ", ".join(p["skills"][:40])]
    L += ["", "---", "_Source: Crustdata people DB (LinkedIn-derived) + public web check of LinkedIn activity, 2026-09-27._"]
    return "\n".join(L) + "\n"


all_raw = []
for role, rows in selected.items():
    d = OUT / "profiles" / role
    d.mkdir(exist_ok=True)
    for i, r in enumerate(rows, 1):
        slug = re.sub(r"[^a-z0-9]+", "-", r["name"].lower()).strip("-")
        r["profile_file"] = f"profiles/{role}/{i:02d}-{slug}.md"
        (OUT / r["profile_file"]).write_text(dossier(r, i))
        all_raw.append({"role": ROLE_LABEL[role], "segment": r["segment"], **raw[r["person_id"]]})
(OUT / "profiles" / "all_profiles_raw.json").write_text(json.dumps(all_raw, indent=1))

# ---------- workbook ----------
F = "Arial"
HDR = PatternFill("solid", start_color="1F3864")
VOICE = PatternFill("solid", start_color="E2EFDA")
OPER = PatternFill("solid", start_color="DDEBF7")
COLS = [("Rank", 6), ("Segment", 26), ("Name", 24), ("Title", 38), ("Company", 32), ("Location", 26),
        ("LinkedIn (vanity)", 42), ("LinkedIn (ID link)", 30), ("Followers", 10), ("Connections", 11),
        ("Yrs exp", 8), ("Yrs at co.", 9), ("LinkedIn activity", 10), ("Activity evidence", 70),
        ("Voice score", 9), ("Work score", 9), ("Verify", 40), ("Profile file", 34)]

wb = Workbook()
ws = wb.active
ws.title = "README"
readme = [
    ("Top LinkedIn profiles - Outpatient / Independent Radiology & Imaging Centers (USA)", True),
    ("Built 2026-09-27 for Gravity by AlphaNodus outreach. Followers shown as n/a where the data source did not report a count.", False),
    ("", False),
    ("What is in each role tab (20 people each):", True),
    ("  Ranks 1-10 = VOICES: visibly active on LinkedIn (authored posts, articles, podcast/speaking, frequent engagement), ranked by Voice score.", False),
    ("  Ranks 11-20 = OPERATORS: top by track record (seniority, title fit, tenure, experience, credentials, network), ranked by Work score.", False),
    ("", False),
    ("Scoring:", True),
    ("  Work score = 2 x seniority level + title bump (C-suite/VP +2, Director +1.2, Manager +0.6) + yrs experience/6 (cap 30) + yrs at company/5 (cap 15) + 0.6 per credential (CRCR, CPC, CHFP, MBA, MHA...) + 0.8 x log10(connections).", False),
    ("  Voice score = 10 (level A) or 5 (level B) + evidence bonus (podcast host, conference speaker, advocacy, very recent posting: 0.5-5) + 1.5 x log10(followers, or connections where followers unavailable).", False),
    ("  LinkedIn activity: A = regularly authors own posts/articles/podcast; B = occasional posts or frequent public comments/engagement; C = profile only, nothing public found.", False),
    ("", False),
    ("Filters applied:", True),
    ("  US-located; current employer name contains Imaging / Radiology / MRI / Diagnostic; employer industry = Hospitals & Health Care or Medical Practices.", False),
    ("  Excluded: hospitals, universities/academic departments, health systems, equipment makers/servicers, AI & teleradiology vendors, billing vendors, labs, mobile-imaging leasing, pain clinics.", False),
    ("  Max 3 people per company per role, so large groups (e.g. Radiology Partners, Lumexa) cannot dominate.", False),
    ("", False),
    ("Honest caveats:", True),
    ("  'Viral' creators are rare in these roles. Revenue-cycle leaders at imaging centers mostly engage (comment) rather than publish; their Voices segment is the most visible of the pool, not influencers.", False),
    ("  Activity was checked via public web search of LinkedIn (posts, articles, comments indexed publicly), not a logged-in LinkedIn scrape, so some private/unindexed activity is missed.", False),
    ("  Follower counts were not pulled for the Revenue Cycle pool; connections are used as the reach proxy there.", False),
    ("  Database snapshot may lag job changes. Rows with a 'Verify' note had a newer headline on the web.", False),
    ("  Contact info (emails/phones) NOT included.", False),
    ("", False),
    ("Tabs: Revenue Cycle | Referral & Growth Lead | Practice Growth | Adjacent Influencers (strong voices who are vendor/coalition-side, not at an imaging center).", False),
    ("Full profile dossiers (work history, education, skills) are in the profiles/ folder next to this file; raw JSON in profiles/all_profiles_raw.json.", False),
]
for i, (t, b) in enumerate(readme, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=b, size=13 if i == 1 else 10)
ws.column_dimensions["A"].width = 160

TAB = {"rcm": "Revenue Cycle", "referral": "Referral & Growth Lead", "growth": "Practice Growth"}
for role, rows in selected.items():
    s = wb.create_sheet(TAB[role])
    for j, (h, w) in enumerate(COLS, 1):
        c = s.cell(row=1, column=j, value=h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
        s.column_dimensions[get_column_letter(j)].width = w
    for i, r in enumerate(rows, 1):
        vals = [i, r["segment"], r["name"], r["title"], r["company"], r["region"], r["vanity_url"] or "",
                r["linkedin_url"], r["followers"] if r["followers"] is not None else "n/a", r["connections"],
                r["years_experience"], r["years_at_company"], r["activity_level"], r["evidence"],
                r["voice_score"], r["work_score"], r["flag"], r["profile_file"]]
        fill = VOICE if r["segment"].startswith("Voice") else OPER
        for j, v in enumerate(vals, 1):
            c = s.cell(row=i + 1, column=j, value=v)
            c.font = Font(name=F, size=10)
            c.fill = fill
            c.alignment = Alignment(wrap_text=j in (4, 14, 17), vertical="top")
            if j in (7, 8) and v:
                c.hyperlink = v
                c.font = Font(name=F, size=10, color="0563C1", underline="single")
    s.freeze_panes = "D2"
    s.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{len(rows) + 1}"

adj = wb.create_sheet("Adjacent Influencers")
adj_rows = [("Nicole Jones-Gerbino, FRBMA", "President, PBS Radiology (radiology RCM/billing company)",
             "https://www.linkedin.com/in/njonesgerbino", activity["Nicole Jones-Gerbino, FRBMA"]["evidence"]),
            ("David Howard, MBA, FRBMA", "Healthcare Marketing Executive, Strategic Radiology (coalition of 50 independent practices)",
             "https://www.linkedin.com/in/david-howard-tx", activity["David Howard, MBA, FRBMA"]["evidence"])]
for j, (h, w) in enumerate([("Name", 28), ("Role", 60), ("LinkedIn", 44), ("Why notable", 90)], 1):
    c = adj.cell(row=1, column=j, value=h)
    c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HDR
    adj.column_dimensions[get_column_letter(j)].width = w
for i, row in enumerate(adj_rows, 2):
    for j, v in enumerate(row, 1):
        c = adj.cell(row=i, column=j, value=v)
        c.font = Font(name=F, size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
adj.cell(row=len(adj_rows) + 3, column=1,
         value="Not imaging-center employees, so excluded from the ranked lists; high-reach voices in the independent-radiology business community, useful for co-marketing or amplification.").font = Font(name=F, italic=True, size=10)

wb.save(OUT / "top_profiles_radiology_imaging.xlsx")
print("saved", file=sys.stderr)
